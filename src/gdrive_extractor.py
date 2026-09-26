"""
透過 Google Drive API（Service Account 驗證）讀取指定資料夾（含所有子資料夾）內的文件，
轉成與 GitHub 來源相同的 ExtractedDocument，供 summarizer.py 轉譯為「功能卡片」。

支援格式：
    - Google 文件 / 簡報（匯出純文字）、Google 試算表（匯出 xlsx，讀取所有工作表）
    - txt / md / csv / tsv / json
    - PDF（僅文字層；掃描影像無法擷取）
    - docx / xlsx / pptx
其餘格式（圖片、影片、表單、捷徑、舊版 .doc/.xls/.ppt 等）一律略過，列入略過清單。
"""
from __future__ import annotations

import io
import json
from dataclasses import dataclass
from typing import Callable, Dict, Iterable, Iterator, List, Optional, Set, Tuple

import docx
import openpyxl
import pptx
import pypdf
from google.auth.exceptions import GoogleAuthError
from google.oauth2 import service_account
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError

from src import config
from src.documents import ExtractedDocument

logger = config.get_logger(__name__)

DRIVE_SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]
FOLDER_MIME = "application/vnd.google-apps.folder"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
PDF_MIME = "application/pdf"
# Google 原生格式 → 匯出格式（試算表匯出 xlsx 才能讀到所有工作表；text/csv 只有第一張）
GOOGLE_EXPORTS = {
    "application/vnd.google-apps.document": "text/plain",
    "application/vnd.google-apps.presentation": "text/plain",
    "application/vnd.google-apps.spreadsheet": XLSX_MIME,
}
TEXT_MIMES = {
    "text/plain",
    "text/markdown",
    "text/x-markdown",
    "text/csv",
    "text/tab-separated-values",
    "application/json",
}
SUPPORTED_MIMES = set(GOOGLE_EXPORTS) | TEXT_MIMES | {PDF_MIME, DOCX_MIME, XLSX_MIME, PPTX_MIME}
MAX_TEXT_CHARS = 40_000  # 與 github_extractor.MAX_FILE_BYTES 同級
MAX_DOWNLOAD_BYTES = 20 * 1024 * 1024
LIST_FIELDS = "nextPageToken, files(id, name, mimeType, modifiedTime, size, webViewLink)"
API_RETRIES = 3


class GDriveExtractorError(Exception):
    """Google Drive 存取相關錯誤，訊息需可直接呈現給使用者。"""


@dataclass
class GDriveExtractResult:
    docs: List[ExtractedDocument]
    seen_file_ids: Set[str]  # 本次掃描到的「支援格式」檔案 ID（含讀取失敗、無文字、過大者）
    skipped: List[str]  # 人類可讀「路徑（原因）」
    complete: bool  # False = 達 max_files 上限或有子資料夾列舉失敗


# ---------------------------------------------------------------------------
# 文字擷取：全部以產生器逐段產出，累積達 MAX_TEXT_CHARS 即停止解析，
# 避免大型試算表 / PDF 全量解析。
# ---------------------------------------------------------------------------
def _join_capped(parts: Iterable[str]) -> str:
    """略過空白片段、以換行串接，累積長度達 MAX_TEXT_CHARS 即停止迭代。"""
    out: List[str] = []
    total = 0
    for part in parts:
        if not part or not part.strip():
            continue
        out.append(part)
        total += len(part) + 1
        if total >= MAX_TEXT_CHARS:
            break
    return "\n".join(out)


def _pdf_to_text(data: bytes) -> str:
    reader = pypdf.PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF 有密碼保護")
    return _join_capped(page.extract_text() or "" for page in reader.pages)


def _docx_lines(document) -> Iterator[str]:
    for paragraph in document.paragraphs:
        yield paragraph.text
    for table in document.tables:
        for row in table.rows:
            yield " | ".join(cell.text.strip() for cell in row.cells)


def _docx_to_text(data: bytes) -> str:
    return _join_capped(_docx_lines(docx.Document(io.BytesIO(data))))


def _xlsx_lines(workbook) -> Iterator[str]:
    for ws in workbook.worksheets:
        header_emitted = False
        for row in ws.iter_rows(values_only=True):
            line = "\t".join("" if v is None else str(v) for v in row).rstrip()
            if not line.strip():
                continue
            if not header_emitted:
                # 工作表標題只在該表有內容時輸出，全空工作表不產生任何文字
                yield f"## 工作表：{ws.title}"
                header_emitted = True
            yield line


def _xlsx_to_text(data: bytes) -> str:
    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return _join_capped(_xlsx_lines(workbook))
    finally:
        workbook.close()


def _pptx_lines(presentation) -> Iterator[str]:
    for idx, slide in enumerate(presentation.slides, start=1):
        texts: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if shape.has_table:
                for row in shape.table.rows:
                    texts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if any(t.strip() for t in texts):
            # 投影片標題只在該頁有文字時輸出，純圖片投影片不產生任何文字
            yield f"## 投影片 {idx}"
            yield from texts


def _pptx_to_text(data: bytes) -> str:
    return _join_capped(_pptx_lines(pptx.Presentation(io.BytesIO(data))))


def _decode_text(data: bytes) -> str:
    # utf-8-sig：Google 文件匯出的純文字與部分上傳的 txt/csv 開頭帶 BOM
    return data.decode("utf-8-sig", errors="replace")


_BINARY_PARSERS: Dict[str, Callable[[bytes], str]] = {
    PDF_MIME: _pdf_to_text,
    DOCX_MIME: _docx_to_text,
    XLSX_MIME: _xlsx_to_text,
    PPTX_MIME: _pptx_to_text,
}


class GDriveExtractor:
    """封裝以 Service Account 讀取 Google Drive 資料夾（含子資料夾）文件的邏輯。"""

    def __init__(
        self,
        credentials_json: Optional[str] = None,
        folder_ids: Optional[List[str]] = None,
        max_files: Optional[int] = None,
    ) -> None:
        raw = (credentials_json or config.GDRIVE_SERVICE_ACCOUNT_JSON).strip()
        self.folder_ids = folder_ids or config.GDRIVE_FOLDER_IDS
        self.max_files = max_files or config.GDRIVE_MAX_FILES

        if not raw:
            raise GDriveExtractorError("GDRIVE_SERVICE_ACCOUNT_JSON 未設定，請在 .env 中設定 Service Account 金鑰。")
        if not self.folder_ids:
            raise GDriveExtractorError("GDRIVE_FOLDER_IDS 未設定，請在 .env 中填入要同步的資料夾 ID。")

        key_path = None if raw.startswith("{") else config.resolve_gdrive_credentials_path(raw)
        try:
            info = json.loads(raw if key_path is None else key_path.read_text(encoding="utf-8"))
            if not isinstance(info, dict):
                raise ValueError("金鑰內容不是 JSON 物件")
            creds = service_account.Credentials.from_service_account_info(info, scopes=DRIVE_SCOPES)
        except FileNotFoundError as exc:
            raise GDriveExtractorError(f"找不到 Service Account 金鑰檔：{key_path}") from exc
        except OSError as exc:
            raise GDriveExtractorError(f"無法讀取 Service Account 金鑰檔 {key_path}：{exc}") from exc
        except (ValueError, KeyError) as exc:  # json.JSONDecodeError 是 ValueError 子類別
            raise GDriveExtractorError(
                f"GDRIVE_SERVICE_ACCOUNT_JSON 不是合法的 Service Account 金鑰 JSON：{exc}"
            ) from exc

        self.service_account_email: str = creds.service_account_email
        try:
            self._service = build("drive", "v3", credentials=creds, cache_discovery=False)
        except Exception as exc:  # noqa: BLE001
            raise GDriveExtractorError(f"初始化 Google Drive 客戶端失敗：{exc}") from exc

    # ------------------------------------------------------------------
    def _get_root(self, folder_id: str) -> dict:
        try:
            root = (
                self._service.files()
                .get(fileId=folder_id, fields="id, name, mimeType, driveId", supportsAllDrives=True)
                .execute(num_retries=API_RETRIES)
            )
        except HttpError as exc:
            status = exc.resp.status
            if status == 404:
                raise GDriveExtractorError(
                    f"找不到 Google Drive 資料夾 `{folder_id}`：請確認 ID 正確，且已將資料夾共用給 "
                    f"Service Account `{self.service_account_email}`（檢視者即可）。"
                ) from exc
            if status == 403:
                raise GDriveExtractorError(
                    f"沒有權限讀取 Google Drive 資料夾 `{folder_id}`：請確認 GCP 專案已啟用 Google Drive API，"
                    f"且資料夾已共用給 `{self.service_account_email}`。（{exc}）"
                ) from exc
            raise GDriveExtractorError(f"連線 Google Drive 失敗：{exc}") from exc
        except GoogleAuthError as exc:
            raise GDriveExtractorError(
                f"無法取得 Google 存取權杖（金鑰無效、已停用，或網路無法連線）：{exc}"
            ) from exc
        except Exception as exc:  # noqa: BLE001
            raise GDriveExtractorError(f"連線 Google Drive 失敗：{exc}") from exc

        mime = root.get("mimeType")
        if mime != FOLDER_MIME:
            raise GDriveExtractorError(
                f"`{folder_id}` 不是資料夾（mimeType={mime}），GDRIVE_FOLDER_IDS 只能填資料夾 ID。"
            )
        return root

    def _list_children(self, folder_id: str, drive_id: Optional[str]) -> List[dict]:
        params = {
            "q": f"'{folder_id}' in parents and trashed = false",
            "fields": LIST_FIELDS,
            "pageSize": 1000,
            "orderBy": "folder,name",
            "supportsAllDrives": True,
            "includeItemsFromAllDrives": True,
        }
        if drive_id:
            params.update(corpora="drive", driveId=drive_id)
        else:
            params["corpora"] = "user"

        items: List[dict] = []
        page_token: Optional[str] = None
        while True:
            response = (
                self._service.files()
                .list(pageToken=page_token, **params)
                .execute(num_retries=API_RETRIES)
            )
            items.extend(response.get("files", []))
            page_token = response.get("nextPageToken")
            if not page_token:
                return items

    def _read_text(self, item: dict) -> str:
        mime = item["mimeType"]
        files = self._service.files()
        if mime in GOOGLE_EXPORTS:
            export_mime = GOOGLE_EXPORTS[mime]
            data = files.export_media(fileId=item["id"], mimeType=export_mime).execute(num_retries=API_RETRIES)
            return _xlsx_to_text(data) if export_mime == XLSX_MIME else _decode_text(data)

        data = files.get_media(fileId=item["id"], supportsAllDrives=True).execute(num_retries=API_RETRIES)
        if mime in TEXT_MIMES:
            return _decode_text(data)
        return _BINARY_PARSERS[mime](data)

    # ------------------------------------------------------------------
    def extract_all(self) -> GDriveExtractResult:
        """深度優先掃描所有設定的資料夾（含子資料夾），讀取支援格式檔案的文字內容。"""
        docs: List[ExtractedDocument] = []
        seen: Set[str] = set()
        skipped: List[str] = []
        visited: Set[str] = set()
        complete = True
        limit_reached = False

        for root_id in self.folder_ids:
            root = self._get_root(root_id)  # 根資料夾錯誤直接 raise
            stack: List[Tuple[str, str, Optional[str]]] = [(root["id"], root["name"], root.get("driveId"))]

            while stack and not limit_reached:
                folder_id, folder_path, drive_id = stack.pop()
                if folder_id in visited:
                    continue
                visited.add(folder_id)

                try:
                    children = self._list_children(folder_id, drive_id)
                except Exception as exc:  # noqa: BLE001
                    logger.warning("列舉 Google Drive 資料夾「%s」失敗，已略過：%s", folder_path, exc)
                    skipped.append(f"{folder_path}/（子資料夾列舉失敗：{exc}）")
                    complete = False
                    continue

                subfolders: List[Tuple[str, str, Optional[str]]] = []
                for item in children:
                    path = f"{folder_path}/{item['name']}"
                    mime = item["mimeType"]
                    if mime == FOLDER_MIME:
                        subfolders.append((item["id"], path, drive_id))
                        continue
                    if mime not in SUPPORTED_MIMES:
                        skipped.append(f"{path}（不支援的格式：{mime}）")
                        continue
                    if item["id"] in seen:  # 多個設定資料夾重疊時去重
                        continue
                    if len(seen) >= self.max_files:
                        limit_reached = True
                        break
                    seen.add(item["id"])

                    if mime not in GOOGLE_EXPORTS and int(item.get("size") or 0) > MAX_DOWNLOAD_BYTES:
                        skipped.append(f"{path}（檔案超過 20MB）")
                        continue
                    try:
                        text = self._read_text(item).strip()
                    except Exception as exc:  # noqa: BLE001
                        logger.warning("讀取 Google Drive 檔案「%s」失敗，已略過：%s", path, exc)
                        skipped.append(f"{path}（讀取失敗：{exc}）")
                        continue
                    if not text:
                        skipped.append(f"{path}（沒有可擷取的文字，可能是掃描影像）")
                        continue

                    docs.append(
                        ExtractedDocument(
                            doc_type="gdrive",
                            title=f"Drive 文件：{path}",
                            content=text[:MAX_TEXT_CHARS],
                            files=[path],
                            metadata={
                                "file_id": item["id"],
                                "mime_type": mime,
                                "modified_time": item.get("modifiedTime"),
                                "url": item.get("webViewLink"),
                            },
                        )
                    )

                stack.extend(reversed(subfolders))  # pop 時依名稱順序

            if limit_reached:
                logger.warning("已達 GDRIVE_MAX_FILES=%d 上限，停止掃描 Google Drive。", self.max_files)
                complete = False
                break

        logger.info("從 Google Drive 讀取到 %d 份文件（略過 %d 份）。", len(docs), len(skipped))
        return GDriveExtractResult(docs=docs, seen_file_ids=seen, skipped=skipped, complete=complete)
